"""
Extracts plain text from uploaded PDF, DOCX, and TXT files.
"""
import io

import fitz  # PyMuPDF
from docx import Document
from fastapi import HTTPException, status

ALLOWED_EXTENSIONS = {"pdf", "txt", "docx"}


def get_file_extension(filename: str) -> str:
    if "." not in filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File must have an extension (.pdf, .txt, or .docx).",
        )
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type '.{ext}'. Allowed types: pdf, txt, docx.",
        )
    return ext


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_chunks = []
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            for page in doc:
                text_chunks.append(page.get_text())
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not read PDF file: {exc}",
        )
    return "\n".join(text_chunks).strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        # Also pull text out of any tables in the document
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        return "\n".join(p for p in paragraphs if p.strip()).strip()
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not read DOCX file: {exc}",
        )


def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8", errors="ignore").strip()
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Could not read TXT file: {exc}",
        )


def extract_text(filename: str, file_bytes: bytes) -> tuple[str, str]:
    """Returns (extracted_text, file_extension)."""
    ext = get_file_extension(filename)

    if ext == "pdf":
        text = extract_text_from_pdf(file_bytes)
    elif ext == "docx":
        text = extract_text_from_docx(file_bytes)
    else:
        text = extract_text_from_txt(file_bytes)

    if not text:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No readable text could be extracted from this file.",
        )

    return text, ext
